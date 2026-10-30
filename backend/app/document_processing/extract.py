"""
Document processing: text extraction, chunking, and storage key generation.
Supports PDF (PyMuPDF), DOCX (python-docx), XLSX (openpyxl), and TXT.
"""
from __future__ import annotations

import hashlib
import io
import logging
import uuid
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

SUPPORTED_MIMES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "text/plain",
}

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt"}


def generate_storage_key(original_filename: str) -> str:
    """Generate a random, safe storage key — never use the client filename as path."""
    ext = Path(original_filename).suffix.lower()
    return f"uploads/{uuid.uuid4().hex}{ext}"


def compute_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def validate_file(filename: str, content_type: str, size_bytes: int, max_mb: int) -> None:
    from app.core.errors import FileTooLargeError, UnsupportedFileTypeError
    if size_bytes > max_mb * 1024 * 1024:
        raise FileTooLargeError(max_mb)
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS and content_type not in SUPPORTED_MIMES:
        raise UnsupportedFileTypeError()


def extract_text(file_bytes: bytes, filename: str, mime_type: str) -> tuple[str, list[dict], int]:
    """
    Extract text from file. Returns (full_text, page_chunks, page_count).
    page_chunks = [{"page": int, "text": str, "section": str|None}]
    """
    ext = Path(filename).suffix.lower()

    if ext == ".pdf" or mime_type == "application/pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes)
    elif ext == ".xlsx":
        return _extract_xlsx(file_bytes)
    elif ext == ".txt":
        return _extract_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(data: bytes) -> tuple[str, list[dict], int]:
    import fitz  # PyMuPDF
    doc = fitz.open(stream=data, filetype="pdf")
    pages = []
    full_parts = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text")
        if text.strip():
            pages.append({"page": page_num + 1, "text": text, "section": None})
            full_parts.append(text)
    return "\n".join(full_parts), pages, len(doc)


def _extract_docx(data: bytes) -> tuple[str, list[dict], int]:
    from docx import Document
    doc = Document(io.BytesIO(data))
    pages = []
    current_section = None
    current_text: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name.lower()
        if "heading" in style_name:
            if current_text:
                pages.append({
                    "page": len(pages) + 1,
                    "text": "\n".join(current_text),
                    "section": current_section,
                })
            current_section = text
            current_text = []
        else:
            current_text.append(text)

    if current_text:
        pages.append({
            "page": len(pages) + 1,
            "text": "\n".join(current_text),
            "section": current_section,
        })

    full_text = "\n".join(p["text"] for p in pages)
    return full_text, pages, len(pages)


def _extract_xlsx(data: bytes) -> tuple[str, list[dict], int]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    pages = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                rows.append(row_text)
        if rows:
            pages.append({
                "page": len(pages) + 1,
                "text": f"Sheet: {sheet_name}\n" + "\n".join(rows),
                "section": sheet_name,
            })
    full_text = "\n\n".join(p["text"] for p in pages)
    return full_text, pages, len(pages)


def _extract_txt(data: bytes) -> tuple[str, list[dict], int]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="replace")
    pages = [{"page": 1, "text": text, "section": None}]
    return text, pages, 1


def chunk_pages(
    pages: list[dict],
    document_id: str,
    workspace_id: str,
    owner_id: str,
    original_filename: str,
    chunk_size: int = 800,
    chunk_overlap: int = 100,
    case_id: Optional[str] = None,
) -> list[dict]:
    """
    Split page texts into overlapping chunks suitable for embedding.
    Returns list of chunk dicts with stable IDs.
    """
    chunks = []
    chunk_idx = 0

    for page_info in pages:
        page_text = page_info["text"]
        page_num = page_info["page"]
        section = page_info.get("section")

        words = page_text.split()
        start = 0
        while start < len(words):
            end = min(start + chunk_size, len(words))
            chunk_text = " ".join(words[start:end])
            if len(chunk_text.strip()) > 50:  # Skip tiny chunks
                stable_id = f"doc-{document_id}-chunk-{chunk_idx}"
                chunks.append({
                    "id": stable_id,
                    "chunk_index": chunk_idx,
                    "text": chunk_text,
                    "page_number": page_num,
                    "section_label": section,
                    "metadata": {
                        "document_id": document_id,
                        "workspace_id": workspace_id,
                        "owner_id": owner_id,
                        "original_filename": original_filename,
                        "page_number": page_num,
                        "section_label": section or "",
                        "case_id": case_id or "",
                        "chunk_index": chunk_idx,
                    },
                })
                chunk_idx += 1
            start += chunk_size - chunk_overlap

    return chunks
